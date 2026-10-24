"""
Tests for the plan -> targets -> weekly check loop and the readers it needs.

Runs on the bundled SYNTHETIC demo exports. With FRESH4U_EXPORTS pointing at the folder of real Priva
exports (never committed), it also checks the July 2026 and week 43 2025 figures we verified by hand.
"""
import os
import shutil
import sys
import tempfile
import unittest

import openpyxl
import pandas as pd

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import run  # noqa: E402
from fresh4u import planner  # noqa: E402
from fresh4u import season_plan as sp  # noqa: E402
from fresh4u import targets as tg  # noqa: E402
from fresh4u import weekly_check as wc  # noqa: E402
from fresh4u.app_data import process_folder  # noqa: E402
from fresh4u.enrich import enrich_performance  # noqa: E402
from fresh4u.priva_reader import ReportError, protect_names  # noqa: E402
from fresh4u.special import read_mgmt_report, read_yield  # noqa: E402
from fresh4u.store import Store  # noqa: E402

DEMO = os.path.join(ROOT, "demo_data")
REAL = os.environ.get("FRESH4U_EXPORTS")


def _run(inbox, *extra):
    d = tempfile.mkdtemp()
    run.main(["--inbox", inbox, "--data", os.path.join(d, "data"), "--out", os.path.join(d, "out"), "--quiet", *extra])
    return d, Store(os.path.join(d, "data"))


class TestDemoLoop(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp, cls.st = _run(DEMO)
        st = cls.st
        cls.y, cls.prod, cls.mr, cls.hrs = (st.history(t) for t in ("yield", "production", "mgmt_report", "hours"))
        cls.rec = enrich_performance(st.history("performance"))
        cls.pace = planner.weekly_pace(cls.rec)
        cls.A = sp.production_anchors(cls.prod, cls.y) + sp.mgmt_anchors(cls.mr, cls.y)
        cls.R = sp.fit_rates(cls.A, cls.hrs)
        cls.pts = sp.crop_work_points(cls.rec, cls.A)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    # ---- pipeline ----
    def test_every_site_report_loads_and_reconciles(self):
        lg = self.st.log()
        loaded = lg[lg["status"] == "loaded"]
        self.assertEqual(set(loaded["report_type"]), {"performance", "production", "hours", "yield", "mgmt_report"})
        rc = pd.read_csv(self.st.recon_path)
        site = rc[rc["file"].str.contains("Production|Hours|Management")]
        self.assertGreater(len(site), 50)
        self.assertTrue(site["match"].all())

    def test_no_names_are_stored(self):
        for t in ("hours", "performance", "production"):
            h = self.st.history(t)
            self.assertFalse(any("name" in c for c in h.columns), t)
            self.assertFalse(h.astype(str).apply(lambda s: s.str.contains("Demo Person")).any().any(), t)
        self.assertTrue(self.hrs["employee_id"].astype(str).str.match(r"^\d+$").all())

    def test_protect_names_uses_numbers_then_pseudonyms(self):
        df = pd.DataFrame({"employee_name": ["Ann", "Ann", "Bo", "Cy"], "employee_number": [7, 8, 9, None]})
        out, notes = protect_names(df)
        self.assertNotIn("employee_name", out)
        self.assertEqual(out["employee_id"].tolist()[:3], ["7", "8", "9"])
        self.assertTrue(out["employee_id"].iloc[3].startswith("P-"))
        self.assertTrue(any("more than one employee number" in n for n in notes))

    def test_yield_with_a_changed_week_is_held_back(self):
        d = tempfile.mkdtemp()
        try:
            src = os.path.join(DEMO, "SYNTHETIC DEMO Production (tactical).xlsx")
            wb = openpyxl.load_workbook(src)
            ws = wb.worksheets[0]
            for r in ws.iter_rows(min_row=8):
                if r[2].value == "Week 20" and r[7].value:
                    r[7].value += 500        # one week's tomato kg, totals left as they were
                    break
            wb.save(os.path.join(d, "edited.xlsx"))
            with self.assertRaises(Exception):
                rep = read_yield(os.path.join(d, "edited.xlsx"))
                self.assertTrue(rep.reconciles)
            _, st = _run(d)
            self.assertEqual(st.log()["status"].tolist(), ["quarantined"])
        finally:
            shutil.rmtree(d, ignore_errors=True)

    def test_broken_management_report_raises(self):
        d = tempfile.mkdtemp()
        try:
            wb = openpyxl.Workbook()
            wb.active.append(["Management Report"])
            for _ in range(6):
                wb.active.append([None])
            p = os.path.join(d, "mr.xlsx")
            wb.save(p)
            with self.assertRaises(ReportError):
                read_mgmt_report(p)
        finally:
            shutil.rmtree(d, ignore_errors=True)

    def test_recursive_folder_skips_demo_and_working_folders(self):
        d = tempfile.mkdtemp()
        try:
            os.makedirs(os.path.join(d, "exports"))
            os.makedirs(os.path.join(d, "repo", "demo_data"))
            shutil.copy(os.path.join(DEMO, "SYNTHETIC DEMO Hours 20260315.xlsx"), os.path.join(d, "exports"))
            shutil.copy(os.path.join(DEMO, "SYNTHETIC DEMO 20260302.xlsx"), os.path.join(d, "repo", "demo_data"))
            res = process_folder(d, os.path.join(d, "store"))
            self.assertEqual(res["log"]["file"].tolist(), ["SYNTHETIC DEMO Hours 20260315.xlsx"])
            again = process_folder(d, os.path.join(d, "store"))     # second run reads nothing new
            self.assertEqual(len(again["log"]), 1)
        finally:
            shutil.rmtree(d, ignore_errors=True)

    # ---- year plan ----
    def test_plan_adds_up(self):
        p = sp.build_plan(sp.weekly_kg(self.y, 2025), self.R, self.pts)
        self.assertEqual(len(p), 52)
        self.assertTrue((p[sp.CATEGORIES] >= 0).all().all())
        self.assertAlmostEqual(float((p[sp.CATEGORIES].sum(axis=1) - p["Total hours"]).abs().max()), 0, places=6)
        hpp = self.R["hours_per_person"].value
        self.assertAlmostEqual(float((p["Total hours"] / hpp - p["People"]).abs().max()), 0, places=6)
        self.assertTrue(p["Changeover"].str.contains("H1–H2").any())

    def test_yield_forecast_moves_harvest_only(self):
        base = sp.build_plan(sp.weekly_kg(self.y, 2025), self.R, self.pts)
        more = sp.build_plan(sp.weekly_kg(self.y, 2025, tom_pct=120), self.R, self.pts)
        self.assertAlmostEqual(more["Tomato harvest"].sum(), 1.2 * base["Tomato harvest"].sum(), places=3)
        self.assertAlmostEqual(more["Strawberry harvest"].sum(), base["Strawberry harvest"].sum(), places=6)

    def test_plan_matches_the_recorded_weeks(self):
        chk = sp.check_against_anchors(self.A, self.R, self.pts, self.y)
        tot = chk[chk["category"] == "Total hours"]
        self.assertEqual(len(tot), 2)
        for r in tot.itertuples():
            self.assertLess(abs(r.plan - r.recorded) / r.recorded, 0.15, r.week)

    def test_every_rate_says_where_it_came_from(self):
        t = sp.rates_table(self.R)
        self.assertTrue(t["evidence"].isin(["Measured", "Measured, thin", "Assumed", "One fortnight"]).all())
        self.assertIn("Assumed", set(t["evidence"]))

    # ---- targets ----
    def test_targets_cover_crop_work_and_production_tasks(self):
        t = tg.build_targets(self.pace, self.prod)
        self.assertTrue({"R&D Analytics registrations", "Production export"} <= set(t["source"]))
        self.assertTrue((t["target_min_per_row"] > 0).all())
        r = t.iloc[0]
        self.assertAlmostEqual(r["priva_norm_m2_per_hour"], 60 * 144.5 / r["target_min_per_row"], places=6)

    def test_a_week_is_never_graded_against_itself(self):
        pace = pd.DataFrame({"week_start": pd.to_datetime(["2026-04-06", "2026-04-13", "2026-04-20"]),
                             "activity_name": ["Clipping"] * 3, "min_per_row": [50.0, 50.0, 80.0],
                             "rows_done": [500.0] * 3, "hours": [416.7, 416.7, 666.7]})
        t = tg.build_targets(pace, None)
        ot = tg.week_on_track(pace, t, "2026-04-20")
        self.assertEqual(ot.loc[0, "target_min_per_row"], 50.0)          # median of the other two weeks
        self.assertEqual(ot.loc[0, "status"], "Slower than target")
        alone = tg.week_on_track(pace.iloc[[2]], t, "2026-04-20")
        self.assertEqual(alone.loc[0, "status"], "No target yet")

    # ---- weekly check ----
    def test_weekly_check_finds_the_planted_slips(self):
        per = wc.periods(self.hrs)[0]
        u = wc.untagged(self.hrs, per)
        self.assertAlmostEqual(u["untagged"].sum(), (u["present"] - u["registered"]).clip(lower=0).sum(), places=6)
        self.assertGreater(int(u["flag"].sum()), 0)
        locs = wc.location_problems(self.prod, per)
        found = set(zip(locs["activity"], locs["problem"]))
        self.assertIn(("Leaf Blowing (S)", "Task for the other crop's house"), found)
        self.assertIn(("Combing (S)", "Hours with no row for a row task"), found)
        lst = wc.supervisor_list(u, wc.slow_registrations(self.rec[self.rec["review_reason"] != ""], per), locs)
        self.assertTrue(lst["hours"].is_monotonic_decreasing)
        self.assertFalse(lst["employee"].astype(str).str.contains("Demo Person").any())


@unittest.skipUnless(REAL and os.path.isdir(REAL), "set FRESH4U_EXPORTS to the folder of real exports")
class TestRealExports(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.d = tempfile.mkdtemp()
        cls.res = process_folder(REAL, os.path.join(cls.d, "data"))

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.d, ignore_errors=True)

    def test_verified_totals(self):
        r = self.res
        self.assertAlmostEqual(r["production"]["hours"].sum(), 12418.18, places=1)
        self.assertAlmostEqual(r["hours"]["registered_hours"].sum(), 12418.18, places=1)
        self.assertAlmostEqual(r["hours"]["paid_hours"].sum(), 11310.0, places=0)
        self.assertAlmostEqual(r["mgmt_report"].query("row_kind == 'block'")["hours"].sum(), 2539.9, places=0)
        self.assertEqual(int((r["hours"].groupby("employee_id")["registered_hours"].sum() > 0).sum()), 136)

    def test_plan_reproduces_both_weeks(self):
        r = self.res
        chk = sp.check_against_anchors(r["anchors"], r["rates"], r["cw_points"], r["yield"])
        for x in chk[chk["category"] == "Total hours"].itertuples():
            self.assertLess(abs(x.plan - x.recorded) / x.recorded, 0.03, x.week)


if __name__ == "__main__":
    unittest.main()
