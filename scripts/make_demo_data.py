"""
Generates SYNTHETIC Priva-style Performance exports for the public demo.

No Fresh4U records are used. Only rough, non-identifying shape parameters
(typical minutes per row per task, crew sizes, 560 rows of 144.5 m2) are
borrowed so the demo behaves like the real thing. Every file is titled
"SYNTHETIC DEMO" and workers are named D001, D002, ...

    python3 scripts/make_demo_data.py demo_data
"""
import os
import shutil
import sys
from datetime import date, timedelta

import numpy as np
import openpyxl

OUT = sys.argv[1] if len(sys.argv) > 1 else "demo_data"
rng = np.random.default_rng(439)

TASKS = {  # name: (typical minutes per full row, crew size)
    "Clipping (T)": (60, 14), "Deleafing (T)": (49, 11), "Truss Clips (T)": (37, 9),
    "Pruning Flowers/Fruit (T)": (23, 5), "Lowering (T)": (20, 4),
}
ROWS = [1000 * h + i for h in (1, 2) for i in range(1, 281)]
ROW_M2, DENSITY = 144.5, 3.3
START = date(2026, 3, 2)          # Monday
WEEKS = 6

# --- people -------------------------------------------------------------------
people, pid = {}, 1
crews = {}
for t, (_, n) in TASKS.items():
    crews[t] = []
    for _ in range(n):
        name = f"D{pid:03d}"
        grp = "TFW" if rng.random() < 0.62 else "Contractor"
        people[name] = dict(group=grp, speed=float(np.exp(rng.normal(0, 0.16))),
                            start_week=0 if rng.random() > 0.15 else int(rng.integers(2, 4)))
        crews[t].append(name)
        pid += 1


def row_factor(r):
    pos = r % 1000
    return 1.12 if pos <= 10 or pos >= 271 else 1.0


def build_registrations():
    regs = []   # (date, task, row, person, end_time, hours)
    for w in range(WEEKS):
        for t, (mpr, _) in TASKS.items():
            crew = [p for p in crews[t] if people[p]["start_week"] <= w]
            days = [START + timedelta(weeks=w, days=d) for d in range(6)]
            order = list(ROWS)
            chunks = np.array_split(order, 6)
            for day, rows in zip(days, chunks):
                clock = {p: 7 * 3600 + int(rng.integers(0, 900)) for p in crew}
                for k, r in enumerate(rows):
                    p = crew[k % len(crew)]
                    ramp = 1.0
                    weeks_in = w - people[p]["start_week"]
                    if people[p]["start_week"] > 0 and weeks_in < 6:
                        ramp = 0.76 + 0.04 * weeks_in
                    m = mpr * row_factor(int(r)) / people[p]["speed"] / ramp * float(np.exp(rng.normal(0, 0.22)))
                    if day.weekday() == 5:
                        m *= 1.08
                    if rng.random() < 0.015:              # forgot to tag out
                        m += float(rng.uniform(90, 180))
                    hours = m / 60
                    clock[p] += int(hours * 3600)
                    end = clock[p]
                    regs.append((day, t, int(r), p, f"{end // 3600:02d}:{(end % 3600) // 60:02d}:{end % 60:02d}", hours))
    return regs


def write_file(path, title, p_start, p_end, regs):
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Sheet")
    W = 18
    def row(**cells):
        r = [None] * W
        for k, v in cells.items():
            r[int(k[1:])] = v
        ws.append(r)
    row(c0=title, c13=f"Period: {p_start} to {p_end}")
    row(c12="Grand Total")
    row(c0="Date", c1="Activity", c4="Location", c6="Employee group", c9="Employee (short)", c10="End time ",
        c12="Registered hours", c14="Result", c15="Area", c16="Performance / hour")

    def vals(hrs, area):
        res = area * DENSITY
        return dict(c12=hrs, c14=res, c15=area, c16=res / hrs if hrs else None)

    grand_h = grand_a = 0.0
    by_day = {}
    for g in regs:
        by_day.setdefault(g[0], []).append(g)
    for day in sorted(by_day):
        dr = by_day[day]
        h, a = sum(x[5] for x in dr), ROW_M2 * len(dr)
        grand_h += h
        grand_a += a
        row(c0=f"{day.isoformat()} Total", **vals(h, a))
        row(c0="Date", c1="Activity", c4="Location", c6="Employee group", c9="Employee (short)", c10="End time ")
        for t in TASKS:
            tr = [x for x in dr if x[1] == t]
            if not tr:
                continue
            row(c2=t, **vals(sum(x[5] for x in tr), ROW_M2 * len(tr)))
            for x in sorted(tr, key=lambda z: z[2]):
                v = vals(x[5], ROW_M2)
                row(c3=f"Row {x[2]}", **v)
                row(c5=people[x[3]]["group"], **v)
                row(c7=x[3], **v)
                row(c8=x[4], **v)
    row(c0="Grand Total", **vals(grand_h, grand_a))
    row(c0="2026-10-23 12:00 PM", c13=" 1 / 1")
    wb.save(path)


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    regs = build_registrations()
    for i in range(0, WEEKS, 2):
        s = START + timedelta(weeks=i)
        e = s + timedelta(days=13)
        sub = [g for g in regs if s <= g[0] <= e]
        name = f"SYNTHETIC DEMO {s.strftime('%Y%m%d')}.xlsx"
        write_file(os.path.join(OUT, name), f"SYNTHETIC DEMO {s.strftime('%Y%m%d')}", s.isoformat(), (e - timedelta(days=1)).isoformat(), sub)
        print("wrote", name, len(sub), "registrations")
    # a tampered copy: one registration's hours changed without fixing Priva's totals
    last = sorted(os.listdir(OUT))[-1]
    wb = openpyxl.load_workbook(os.path.join(OUT, last))
    ws = wb.worksheets[0]
    for r in ws.iter_rows(min_row=10):
        if r[8].value and r[12].value:          # an end-time (leaf) row
            r[12].value = r[12].value + 5
            break
    tampered = last.replace(".xlsx", " (edited by hand).xlsx")
    wb.save(os.path.join(OUT, tampered))
    print("wrote", tampered, "(should be quarantined)")
